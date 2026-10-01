from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image
import os
NAVY="152B56"; GREEN="166257"; LEAF="1F8F74"; GREY="556070"; RULE="D9DEE7"
BUILD=os.path.dirname(os.path.abspath(__file__))
def rgb(h): return RGBColor.from_string(h)

def clear(slide, keep_ids):
    for sh in list(slide.shapes):
        if sh.shape_id not in keep_ids:
            sh._element.getparent().remove(sh._element)

def rect(slide,x,y,w,h,fill=None,line=None,lw=0.75,shape=MSO_SHAPE.RECTANGLE):
    s=slide.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    s.shadow.inherit=False
    if fill: s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb=rgb(line); s.line.width=Pt(lw)
    else: s.line.fill.background()
    return s

def text(slide,x,y,w,h,paras,anchor=MSO_ANCHOR.TOP,align=PP_ALIGN.LEFT):
    """paras: list of list-of-runs; run=(text,font,size,color,bold) ; para may be dict(runs=,space=)"""
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=anchor
    for i,pa in enumerate(paras):
        runs=pa["runs"] if isinstance(pa,dict) else pa
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        if isinstance(pa,dict):
            if "space" in pa: p.space_before=Pt(pa["space"])
            if "line" in pa: p.line_spacing=pa["line"]
        for t,font,size,color,bold in runs:
            r=p.add_run(); r.text=t; r.font.name=font; r.font.size=Pt(size)
            r.font.color.rgb=rgb(color); r.font.bold=bold
    return tb

def crop(path,aspect,cx=0.5,cy=0.5,zoom=1.0):
    im=Image.open(path).convert("RGB"); W,H=im.size
    if W/H>aspect: h=H/zoom; w=h*aspect
    else: w=W/zoom; h=w/aspect
    x=min(max(cx*W-w/2,0),W-w); y=min(max(cy*H-h/2,0),H-h)
    return im.crop((int(x),int(y),int(x+w),int(y+h)))

def save(im,name,minw=0):
    if im.width<minw: im=im.resize((minw,int(im.height*minw/im.width)),Image.LANCZOS)
    p=os.path.join(BUILD,name); im.save(p,quality=92); return p

def gradient_lower(im,y0f=0.2,y1f=0.58,peak=0.93):
    """bake a navy ramp into the lower part of the photo: clear above y0f, peak by y1f"""
    im=im.convert("RGBA"); W,H=im.size
    col=Image.new("L",(1,H))
    for y in range(H):
        t=(y/H-y0f)/(y1f-y0f); t=min(max(t,0),1)
        col.putpixel((0,y),int(255*peak*(t*t*(3-2*t))))
    a=col.resize((W,H))
    ov=Image.new("RGBA",(W,H),(0x15,0x2B,0x56,255)); ov.putalpha(a)
    im.alpha_composite(ov); return im.convert("RGB")

def shp(slide,sid):
    for sh in slide.shapes:
        if sh.shape_id==sid: return sh
    raise KeyError(sid)

def put_picture(slide,old_id,path,x,y,w,h):
    """insert picture at the z-position of an existing shape, which is removed"""
    old=shp(slide,old_id)
    pic=slide.shapes.add_picture(path,Inches(x),Inches(y),Inches(w),Inches(h))
    old._element.addprevious(pic._element); old._element.getparent().remove(old._element)
    return pic

def to_back_of(slide,new,ref_id):
    """move shape `new` just before shape ref_id in z-order"""
    shp(slide,ref_id)._element.addprevious(new._element)

def remove(slide,*ids):
    for i in ids:
        e=shp(slide,i)._element; e.getparent().remove(e)

def geom(slide,sid,x=None,y=None,w=None,h=None):
    s=shp(slide,sid)
    if x is not None: s.left=Inches(x)
    if y is not None: s.top=Inches(y)
    if w is not None: s.width=Inches(w)
    if h is not None: s.height=Inches(h)
    return s

def scale_font(shape,f):
    for el in shape._element.iter():
        if el.tag.endswith('}defRPr') or el.tag.endswith('}rPr') or el.tag.endswith('}endParaRPr'):
            if el.get('sz'): el.set('sz',str(int(int(el.get('sz'))*f)))

def set_color(shape,hexv):
    ns='{http://schemas.openxmlformats.org/drawingml/2006/main}'
    for c in shape._element.iter(ns+'srgbClr'): c.set('val',hexv)

def nofill(shape):
    shape.fill.background()
